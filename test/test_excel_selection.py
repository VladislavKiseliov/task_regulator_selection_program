import unittest
from pathlib import Path

# Ensure project root on sys.path for direct runs
import sys
ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.utils.ExelMethod import ExelMethod

try:
    import openpyxl
except Exception:  # pragma: no cover - skip if dependency is missing
    openpyxl = None


class TestExcelSelection(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if openpyxl is None:
            raise unittest.SkipTest("openpyxl not installed")

        xlsx = cls._find_workbook()
        if xlsx is None:
            raise unittest.SkipTest("xlsx file not found in project root")

        cls.workbook = openpyxl.load_workbook(xlsx, data_only=True)
        cls.excel = ExelMethod()
        cls.max_cell_value = cls._find_max_cell_value(cls.workbook)

    @staticmethod
    def _find_workbook():
        # Prefer a file with known name, otherwise pick the first .xlsx in root
        root = Path(".")
        candidates = list(root.glob("*.xlsx"))
        for p in candidates:
            if "Регуляторы давления" in p.name:
                return p
        return candidates[0] if candidates else None

    @staticmethod
    def _parse_header(value):
        if value is None:
            return None
        s = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".")
        parts = s.split("-")
        try:
            if len(parts) == 2:
                v1 = float(parts[0])
                v2 = float(parts[1])
                return {"min": min(v1, v2), "max": max(v1, v2), "is_range": True}
            v = float(parts[0])
            return {"min": v, "max": v, "is_range": False}
        except ValueError:
            return None

    @classmethod
    def _find_sample(cls, require_range=False):
        wb = cls.workbook
        for name in wb.sheetnames:
            sh = wb[name]
            for r in range(4, sh.max_row + 1):
                inlet_h = cls._parse_header(sh.cell(r, 1).value)
                if inlet_h is None:
                    continue
                for c in range(2, sh.max_column + 1):
                    val = sh.cell(r, c).value
                    if not isinstance(val, (int, float)):
                        continue
                    out_h = cls._parse_header(sh.cell(3, c).value)
                    if out_h is None:
                        continue
                    if require_range and not (inlet_h["is_range"] or out_h["is_range"]):
                        continue
                    inlet = (inlet_h["min"] + inlet_h["max"]) / 2
                    outlet = (out_h["min"] + out_h["max"]) / 2
                    bandwidth = int(val)
                    return inlet, outlet, bandwidth
        return None

    @staticmethod
    def _find_max_cell_value(workbook):
        max_val = 0
        for name in workbook.sheetnames:
            sh = workbook[name]
            for row in sh.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, (int, float)) and v > max_val:
                        max_val = v
        return max_val

    def test_conduct_analysis_match(self):
        sample = self._find_sample()
        if sample is None:
            self.skipTest("no suitable numeric sample found in workbook")
        inlet, outlet, bandwidth = sample
        found = self.excel.conduct_analysis(
            self.workbook, inlet, outlet, bandwidth, load_range=(0.2, 1.0)
        )
        self.assertTrue(found, "Expected at least one regulator match")

    def test_conduct_analysis_range_match(self):
        sample = self._find_sample(require_range=True)
        if sample is None:
            self.skipTest("no range-based sample found in workbook")
        inlet, outlet, bandwidth = sample
        found = self.excel.conduct_analysis(
            self.workbook, inlet, outlet, bandwidth, load_range=(0.2, 1.0)
        )
        self.assertTrue(found, "Expected at least one regulator match for range headers")

    def test_conduct_analysis_no_match(self):
        if self.max_cell_value == 0:
            self.skipTest("no numeric values in workbook")
        bandwidth = int(self.max_cell_value * 2 + 1)
        # Use any valid pressures from a sample to avoid early failures
        sample = self._find_sample()
        if sample is None:
            self.skipTest("no suitable numeric sample found in workbook")
        inlet, outlet, _ = sample
        found = self.excel.conduct_analysis(
            self.workbook, inlet, outlet, bandwidth, load_range=(0.2, 1.0)
        )
        self.assertFalse(found, "Expected no regulators for oversized bandwidth")


if __name__ == "__main__":
    unittest.main()
